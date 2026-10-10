"""Build the editable usage handbook from tutorial prose and live package examples."""

from __future__ import annotations

import argparse
import ast
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "usage-handbook.md"
OUTPUT = ROOT / "docs" / "TunaRAG_Python_Usage_Handbook.docx"


def expand_source() -> str:
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    sections = {}
    for match in re.finditer(r"^## ([^\n]+)\n(.*?)(?=^## |\Z)", readme, re.M | re.S):
        sections[match.group(1)] = match.group(2).strip()

    def feature(match):
        body = sections[match.group(1)]
        body = re.sub(r"^### ", "#### ", body, flags=re.M)
        body = body.replace("V0.2", "V0.3.1")
        return body

    def example(match):
        path = ROOT / "examples" / match.group(1)
        content = path.read_text(encoding="utf-8")
        if path.name == "langchain_full_rag.py":
            content = content.replace(
                "Path(__file__).resolve().parents[1]", "Path(__file__).resolve().parent"
            )
            content = content.replace(".env in the repository root", ".env beside this script")
            content = content.replace("repository-root .env", "script-folder .env")
            content = content.replace(
                "python examples/langchain_full_rag.py", "python langchain_full_rag.py"
            )
        return "```python\n" + content.rstrip() + "\n```"

    content = SOURCE.read_text(encoding="utf-8")
    content = re.sub(r"\{\{feature:([^}]+)\}\}", feature, content)
    content = re.sub(r"\{\{example:([^}]+)\}\}", example, content)
    return content


def python_blocks(content: str) -> list[str]:
    return re.findall(r"^```python\n(.*?)^```", content, re.M | re.S)


def validate(content: str) -> None:
    for number, block in enumerate(python_blocks(content), start=1):
        ast.parse(block, filename=f"handbook-example-{number}")
    assert len(python_blocks(content)) >= 40
    assert len(re.findall(r"^## Chapter ", content, re.M)) == 23
    assert "{{feature:" not in content and "{{example:" not in content
    assert "No paid provider calls" in content


def build(output: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    content = expand_source()
    validate(content)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    section.header_distance = section.footer_distance = Inches(0.3)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.06
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Title"].font.size = Pt(30)
    document.styles["Heading 1"].font.size = Pt(20)
    document.styles["Heading 1"].paragraph_format.page_break_before = True
    document.styles["Heading 2"].font.size = Pt(14)
    document.styles["Heading 3"].font.size = Pt(12)
    document.styles["Heading 4"].font.size = Pt(12)
    code_style = document.styles.add_style("Handbook Code", 1)
    code_style.font.name, code_style.font.size = "Consolas", Pt(8.5)
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1
    code_style.paragraph_format.left_indent = Inches(0.1)
    code_style.paragraph_format.right_indent = Inches(0.1)
    code_style.paragraph_format.widow_control = False
    header = section.header.paragraphs[0]
    header.text = "TunaRAG Python Usage Handbook  |  Version 0.3.1"
    header.style = document.styles["Normal"]
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    def formatted(paragraph, text):
        text = re.sub(r"\[([^]]+)\]\(#[^)]+\)", r"\1", text)
        text = re.sub(r"\[([^]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
        for part in re.split(r"(`[^`]+`|\*\*.+?\*\*)", text):
            if not part:
                continue
            run = paragraph.add_run(part.strip("`*") if part.startswith(("`", "**")) else part)
            if part.startswith("`"):
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            elif part.startswith("**"):
                run.bold = True

    def table(rows):
        cells = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows]
        cells = [row for row in cells if not all(re.fullmatch(r"[: -]+", cell) for cell in row)]
        result = document.add_table(rows=0, cols=max(map(len, cells)))
        result.style = "Table Grid"
        for row_number, values in enumerate(cells):
            row = result.add_row()
            properties = row._tr.get_or_add_trPr()
            no_split = OxmlElement("w:cantSplit")
            properties.append(no_split)
            if row_number == 0:
                repeat = OxmlElement("w:tblHeader")
                properties.append(repeat)
            for cell, value in zip(row.cells, values):
                paragraph = cell.paragraphs[0]
                formatted(paragraph, value)
                paragraph.paragraph_format.space_after = Pt(4)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if row_number == 0:
                        run.bold = True
                if row_number == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E7EDF2")
                    cell._tc.get_or_add_tcPr().append(shade)
        document.add_paragraph().paragraph_format.space_after = Pt(1)

    lines = content.splitlines()
    index = 0
    code_number = 0
    chapter_titles = re.findall(r"^## (Chapter [^\n]+|Appendix [^\n]+)$", content, re.M)
    bookmarks = {title: f"handbook_{number}" for number, title in enumerate(chapter_titles)}
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            language = line[3:] or "text"
            block = []
            index += 1
            while index < len(lines) and lines[index] != "```":
                block.append(lines[index])
                index += 1
            code_number += 1
            caption = document.add_paragraph(f"Example {code_number}  |  {language}")
            caption.paragraph_format.keep_with_next = True
            caption.runs[0].bold = True
            caption.runs[0].font.size = Pt(9)
            for line_number, code_line in enumerate(block):
                paragraph = document.add_paragraph(code_line, "Handbook Code")
                paragraph.paragraph_format.keep_with_next = line_number < 2 or (
                    len(block) - 4 <= line_number < len(block) - 1
                )
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "F3F5F7")
                paragraph._p.get_or_add_pPr().append(shade)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
        elif line == "{{contents}}":
            document.add_heading("Contents", 1)
            for title in chapter_titles:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(3)
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(qn("w:anchor"), bookmarks[title])
                run = OxmlElement("w:r")
                text = OxmlElement("w:t")
                text.text = title
                run.append(text)
                hyperlink.append(run)
                paragraph._p.append(hyperlink)
            document.add_paragraph(
                "Use the Navigation pane to jump between chapters. "
                "The entries above link directly to the chapter headings."
            )
        elif line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                rows.append(lines[index])
                index += 1
            table(rows)
            continue
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
            if match:
                relative = match.group(2).split("/assets/", 1)[-1]
                path = ROOT / "docs" / "assets" / relative
                if path.exists():
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.keep_with_next = True
                    paragraph.add_run().add_picture(str(path), width=Inches(6.6))
                    caption = document.add_paragraph(match.group(1))
                    caption.style = "Caption"
        elif line.startswith("# "):
            document.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            title = line[3:]
            heading = document.add_heading(title, 1)
            if title in bookmarks:
                number = chapter_titles.index(title)
                start = OxmlElement("w:bookmarkStart")
                start.set(qn("w:id"), str(number))
                start.set(qn("w:name"), bookmarks[title])
                end = OxmlElement("w:bookmarkEnd")
                end.set(qn("w:id"), str(number))
                heading._p.insert(1, start)
                heading._p.append(end)
        elif line.startswith("### "):
            document.add_heading(line[4:], 2)
        elif line.startswith("#### "):
            document.add_heading(line[5:], 3)
        elif line.startswith(("- ", "* ")):
            formatted(document.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"^\d+\. ", line):
            formatted(document.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
        elif line.strip() and line.strip() != "---":
            paragraph_lines = [line.removeprefix("> ")]
            while (
                index + 1 < len(lines)
                and lines[index + 1].strip()
                and not re.match(r"^(#|```|\||!\[|- |\* |\d+\. |\{\{)", lines[index + 1])
            ):
                index += 1
                paragraph_lines.append(lines[index].removeprefix("> "))
            formatted(document.add_paragraph(), " ".join(paragraph_lines))
        index += 1

    document.core_properties.title = "TunaRAG Python Usage Handbook"
    document.core_properties.subject = "Practical progressive usage of TunaRAG Python 0.3.1"
    document.core_properties.author = "TunaRAG Documentation"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(
        f"Created {output}; {len(python_blocks(content))} Python blocks; "
        f"{code_number} total code examples"
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--check", action="store_true")
    arguments = parser.parse_args()
    if arguments.check:
        validate(expand_source())
        print("Handbook source and Python syntax checks passed")
    else:
        build(arguments.output)


if __name__ == "__main__":
    main()
